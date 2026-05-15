"""
Unit tests for the DAM Word generator (services/sil_proposta/export_service.py).

Parses the generated .docx with python-docx and verifies key content and
structure:
  - Cover page carries the proposal title.
  - All 8 numbered sections present.
  - Premissas + deliverables + resources flow through correctly.
  - Investment value formats as Brazilian BRL.
  - dam_json values override Proposal columns when both present.
"""
from types import SimpleNamespace

import pytest
from docx import Document

from services.sil_proposta.export_service import (
    _format_brl,
    _proposal_to_sections,
    generate_dam_document,
)

pytestmark = pytest.mark.unit


# ── stubs that look enough like SQLAlchemy models for the renderer ──────────


def _resource(**overrides):
    base = {"frente": "SD", "nivel": "Senior", "dias": 10, "horas": 80}
    base.update(overrides)
    return SimpleNamespace(**base)


def _deliverable(**overrides):
    base = {"module": "SD", "item": "Especificação funcional"}
    base.update(overrides)
    return SimpleNamespace(**base)


def _premise(text: str):
    return SimpleNamespace(text=text)


def _proposal(
    *,
    title="Reforma Tributária — Cenário 1",
    project_type="Adequação Fiscal",
    sap_version="ECC 6.0 EHP8",
    states=("SP", "RJ"),
    rfp_text="Cliente precisa adequar emissão fiscal das máquinas POS.",
    needs_cpi=False,
    total_hours=528,
    valor=121_440.00,
    dam_json=None,
    resources=None,
    deliverables=None,
    premises=None,
):
    if resources is None:
        resources = [_resource(frente="SD", dias=5, horas=40), _resource(frente="FI", dias=8, horas=64)]
    if deliverables is None:
        deliverables = [_deliverable(module="SD", item="BAPI Z"), _deliverable(module="FI", item="Trigger evento")]
    if premises is None:
        premises = [_premise("Acessos liberados antes do kickoff."), _premise("Janela QAS disponível.")]
    dam = SimpleNamespace(dam_json=dam_json) if dam_json is not None else None
    return SimpleNamespace(
        title=title,
        project_type=project_type,
        sap_version=sap_version,
        states=list(states),
        rfp_text=rfp_text,
        needs_cpi=needs_cpi,
        total_hours=total_hours,
        valor=valor,
        dam=dam,
        resources=resources,
        deliverables=deliverables,
        premises=premises,
    )


def _doc_text(buf) -> str:
    """Concatenated text of paragraphs + table cells from the generated docx."""
    doc = Document(buf)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


# ── adapter ─────────────────────────────────────────────────────────────────


def test_proposal_to_sections_pulls_from_columns_when_dam_json_absent():
    sections = _proposal_to_sections(_proposal())
    assert sections["titulo"] == "Reforma Tributária — Cenário 1"
    assert sections["versao_sap"] == "ECC 6.0 EHP8"
    assert sections["ufs"] == ["SP", "RJ"]
    assert sections["necessidade"].startswith("Cliente precisa adequar")
    assert sections["entregaveis"] == ["[SD] BAPI Z", "[FI] Trigger evento"]
    assert sections["premissas"] == ["Acessos liberados antes do kickoff.", "Janela QAS disponível."]
    assert sections["total_horas"] == 528
    assert sections["comercial"]["valor_referencia"] == 121_440.00


def test_proposal_to_sections_dam_json_overrides_columns():
    dam_json = {
        "titulo": "DAM Gerado pela IA",
        "necessidade": "Texto refinado pelo agente Orion.",
        "comercial": {
            "valor_referencia": 200_000.00,
            "faturamento": "30/30/40",
            "garantia": "60 dias",
            "validade": "45 dias",
        },
    }
    sections = _proposal_to_sections(_proposal(dam_json=dam_json))
    assert sections["titulo"] == "DAM Gerado pela IA"
    assert sections["necessidade"] == "Texto refinado pelo agente Orion."
    assert sections["comercial"]["valor_referencia"] == 200_000.00
    assert sections["comercial"]["faturamento"] == "30/30/40"


# ── renderer ────────────────────────────────────────────────────────────────


def test_generate_dam_produces_nonempty_docx():
    buf = generate_dam_document(_proposal())
    assert buf.getbuffer().nbytes > 0
    # Docx files are ZIPs; the magic bytes are "PK".
    buf.seek(0)
    assert buf.read(2) == b"PK"


def test_generate_dam_cover_carries_title():
    buf = generate_dam_document(_proposal(title="Adequação NF-e Cenário Maquininha"))
    text = _doc_text(buf)
    assert "Adequação NF-e Cenário Maquininha" in text
    assert "CAST GROUP PARTNER" in text
    assert "DAM — Documento de Arquitetura de Melhoria — v1" in text


def test_generate_dam_has_all_eight_sections():
    buf = generate_dam_document(_proposal())
    text = _doc_text(buf)
    for heading in (
        "1. Necessidade",
        "2. Solução",
        "3. Premissas Gerais",
        "4. Equipe do Projeto",
        "5. Análise de Impactos",
        "6. Cronograma",
        "7. Investimento",
        "8. Condições de Faturamento",
    ):
        assert heading in text, f"missing section heading: {heading}"


def test_generate_dam_lists_premissas():
    buf = generate_dam_document(
        _proposal(
            premises=[
                _premise("Premissa única e específica do cenário."),
                _premise("Outra premissa que precisa aparecer."),
            ]
        )
    )
    text = _doc_text(buf)
    assert "Premissa única e específica do cenário." in text
    assert "Outra premissa que precisa aparecer." in text


def test_generate_dam_lists_equipe():
    buf = generate_dam_document(
        _proposal(
            resources=[
                _resource(frente="ABAP", nivel="Pleno"),
                _resource(frente="SD", nivel="Sênior"),
            ]
        )
    )
    text = _doc_text(buf)
    assert "Consultor ABAP" in text
    assert "Pleno" in text
    assert "Consultor SD" in text


def test_generate_dam_formats_brl_value():
    buf = generate_dam_document(_proposal(valor=121_440.00))
    text = _doc_text(buf)
    assert "R$ 121.440,00" in text


def test_generate_dam_uses_dam_json_value_when_present():
    buf = generate_dam_document(
        _proposal(
            valor=121_440.00,
            dam_json={"comercial": {"valor_referencia": 250_000.00}},
        )
    )
    text = _doc_text(buf)
    assert "R$ 250.000,00" in text


def test_generate_dam_needs_cpi_routing_phrase():
    buf_no_cpi = generate_dam_document(_proposal(needs_cpi=False))
    buf_cpi = generate_dam_document(_proposal(needs_cpi=True))
    text_no_cpi = _doc_text(buf_no_cpi)
    text_cpi = _doc_text(buf_cpi)
    assert "SAP DRC" in text_no_cpi
    assert "SAP CPI" in text_cpi


def test_generate_dam_handles_empty_proposal_gracefully():
    # No resources, no deliverables, no premises, no rfp_text.
    buf = generate_dam_document(
        _proposal(
            resources=[],
            deliverables=[],
            premises=[],
            rfp_text=None,
            valor=None,
            total_hours=None,
        )
    )
    assert buf.getbuffer().nbytes > 0
    text = _doc_text(buf)
    assert "1. Necessidade" in text


# ── BRL helper ──────────────────────────────────────────────────────────────


def test_format_brl():
    assert _format_brl(0) == "R$ 0,00"
    assert _format_brl(1234.5) == "R$ 1.234,50"
    assert _format_brl(1_000_000) == "R$ 1.000.000,00"
