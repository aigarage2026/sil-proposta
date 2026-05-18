"""
Document text extractors.

Each loader returns the plain text of a file with all embedded media
(images, logos) stripped — extraction targets *content*, not layout. The
PDF loader uses PyMuPDF (fitz) and skips image blocks entirely. The DOCX
loader walks paragraphs + table cells. The XLSX loader concatenates
non-empty cell values row-by-row.

.doc (legacy Word binary) is intentionally NOT handled here. The CLI
shells out to `libreoffice --headless --convert-to docx` first, then
hands the resulting .docx back through `load_document()`.

Why no image extraction at all: the historical corpus contains client
logos in headers/footers. Even if we anonymize text, leaving images
behind defeats the purpose. Strip-on-read is the simplest guarantee.
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional


def load_document(path: Path) -> Optional[str]:
    """Dispatch on extension. Returns plain text or None for unsupported."""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".xlsx":
        return _load_xlsx(path)
    return None


def _load_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _load_pdf(path: Path) -> str:
    import fitz  # PyMuPDF

    parts: list[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            # get_text("text") emits text blocks only; images are skipped.
            text = (page.get_text("text") or "").strip()
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def _load_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        sheet_lines: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                sheet_lines.append(" | ".join(cells))
        if sheet_lines:
            parts.append(f"[Sheet: {ws.title}]")
            parts.extend(sheet_lines)
    wb.close()
    return "\n".join(parts)
