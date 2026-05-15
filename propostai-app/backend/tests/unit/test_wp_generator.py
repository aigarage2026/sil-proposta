"""
Unit tests for the WP Excel generator (services/propostai/export_service.py).

Loads the produced .xlsx with openpyxl and checks structure + totals:
  - Workbook has one sheet "WP_RFP".
  - Title row, week headers, resource rows, KT AMS row, grand-total row present.
  - Per-resource week distribution matches the routing rule (SD/FI start in
    week 1, ABAP in week 2, GP one day per week).
  - TOTAL GERAL = sum of resource hours + 16 (KT AMS).
  - Pre-sale (hours_presale) info row appears only when non-zero.
"""
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from services.propostai.export_service import (
    HORAS_DIA,
    SEMANAS_MAX,
    _render_wp,
    _resource_week_distribution,
    generate_wp_workbook,
)

pytestmark = pytest.mark.unit


# ── stubs ──────────────────────────────────────────────────────────────────


def _r(frente, nivel="Senior", dias=5, horas=None):
    return SimpleNamespace(frente=frente, nivel=nivel, dias=dias, horas=horas or dias * HORAS_DIA)


def _proposal(*, resources=None, hours_presale=0):
    if resources is None:
        resources = [_r("SD", dias=5), _r("FI", dias=10), _r("ABAP 1", dias=8)]
    return SimpleNamespace(resources=resources, hours_presale=hours_presale)


def _all_cells_text(ws) -> list[str]:
    """Flattened list of stringified non-empty cell values."""
    out = []
    for row in ws.iter_rows(values_only=True):
        for v in row:
            if v is not None and v != "":
                out.append(str(v))
    return out


# ── distribution helper ─────────────────────────────────────────────────────


def test_sd_resource_fills_from_week_1():
    dist = _resource_week_distribution("SD", 10)
    assert dist == {1: 5, 2: 5}


def test_fi_resource_overflows_to_week_3():
    dist = _resource_week_distribution("FI", 13)
    assert dist == {1: 5, 2: 5, 3: 3}


def test_abap_resource_starts_in_week_2():
    dist = _resource_week_distribution("ABAP 1", 5)
    assert dist == {2: 5}


def test_abap_long_resource_walks_forward():
    dist = _resource_week_distribution("ABAP 2", 14)
    assert dist == {2: 5, 3: 5, 4: 4}


def test_gp_distributes_one_day_per_week():
    dist = _resource_week_distribution("GP", 4)
    assert dist == {1: 1, 2: 1, 3: 1, 4: 1}


def test_distribution_capped_at_semanas_max():
    # 50 days clipped to SEMANAS_MAX weeks × 5 days = 30 max.
    dist = _resource_week_distribution("SD", 50)
    assert max(dist.keys()) == SEMANAS_MAX
    assert sum(dist.values()) == 5 * SEMANAS_MAX


# ── workbook shape ──────────────────────────────────────────────────────────


def test_workbook_has_expected_sheet_and_title():
    buf = generate_wp_workbook(_proposal())
    wb = load_workbook(buf)
    assert wb.sheetnames == ["WP_RFP"]
    ws = wb["WP_RFP"]
    assert ws.cell(row=1, column=1).value == "WORK PACKAGE — PROPOSTAI"


def test_workbook_carries_week_headers_and_totals():
    buf = generate_wp_workbook(_proposal())
    ws = load_workbook(buf)["WP_RFP"]
    # header row 3
    assert ws.cell(row=3, column=1).value == "Frente"
    assert ws.cell(row=3, column=2).value == "Recurso"
    assert ws.cell(row=3, column=3).value == "Nível"
    for s in range(1, SEMANAS_MAX + 1):
        assert ws.cell(row=3, column=3 + s).value == f"Sem {s}"
    # totals
    col_td = 3 + SEMANAS_MAX + 1
    col_th = col_td + 1
    assert ws.cell(row=3, column=col_td).value == "Total Dias"
    assert ws.cell(row=3, column=col_th).value == "Total Horas"


def test_workbook_lists_each_resource():
    buf = generate_wp_workbook(_proposal(
        resources=[_r("SD", dias=5), _r("FI", dias=7), _r("ABAP 1", dias=10)]
    ))
    ws = load_workbook(buf)["WP_RFP"]
    text = " ".join(_all_cells_text(ws))
    assert "SD" in text
    assert "FI" in text
    assert "ABAP 1" in text


def test_workbook_includes_kt_ams_row():
    buf = generate_wp_workbook(_proposal())
    ws = load_workbook(buf)["WP_RFP"]
    text = " ".join(_all_cells_text(ws))
    assert "KT AMS" in text
    assert "Deploy — KT AMS (obrigatório)" in text


def test_total_geral_sums_resource_hours_plus_kt_ams():
    resources = [_r("SD", dias=5), _r("FI", dias=5), _r("ABAP 1", dias=5)]
    buf = generate_wp_workbook(_proposal(resources=resources))
    ws = load_workbook(buf)["WP_RFP"]

    # Distribution sums to total_dias = 5 + 5 + 5 = 15; horas = 15 * 8 = 120;
    # + KT AMS 16 = 136.
    expected_total = 5 * HORAS_DIA * 3 + 16

    # Grand total is in the rightmost column of the "TOTAL GERAL" row. We
    # locate that row by its label.
    col_th = 3 + SEMANAS_MAX + 2
    grand_row = None
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "TOTAL GERAL DO PROJETO":
            grand_row = row
            break
    assert grand_row is not None, "grand total row not found"
    assert ws.cell(row=grand_row, column=col_th).value == expected_total


def test_presale_info_row_only_when_nonzero():
    # zero pre-sale → no internal-cost block
    buf = generate_wp_workbook(_proposal(hours_presale=0))
    text_no_presale = " ".join(_all_cells_text(load_workbook(buf)["WP_RFP"]))
    assert "CUSTO INTERNO" not in text_no_presale

    # non-zero → block appears
    buf = generate_wp_workbook(_proposal(hours_presale=24))
    text_with_presale = " ".join(_all_cells_text(load_workbook(buf)["WP_RFP"]))
    assert "CUSTO INTERNO" in text_with_presale
    assert "24h" in text_with_presale


def test_empty_resources_falls_back_to_sample_team():
    # Empty resources list → renderer fills with the sample team; document
    # is still usable (this guards drafts that haven't been generated).
    buf = _render_wp(resources=[], presale_hours=0)
    ws = load_workbook(buf)["WP_RFP"]
    text = " ".join(_all_cells_text(ws))
    # Sample team includes ABAP 1/2/3
    assert "ABAP 1" in text
    assert "ABAP 2" in text
    assert "ABAP 3" in text
